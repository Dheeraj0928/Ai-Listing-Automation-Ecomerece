"""Bulk operations service — CSV/Excel import, export, bulk AI generation."""

import csv
import io
import json
import logging
import uuid
from datetime import datetime, timezone

import redis.asyncio as aioredis
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.exceptions import ValidationError
from app.models.product import Product
from app.repositories.product_repo import ProductRepository

logger = logging.getLogger(__name__)

# CSV template columns
TEMPLATE_COLUMNS = [
    "sku", "product_name", "brand", "product_type", "description",
    "short_description", "price", "mrp", "cost_price", "stock",
    "color", "size", "material", "weight", "country_of_origin",
    "manufacturer", "manufacturer_address", "packer", "importer",
    "status",
]


def _get_redis():
    return aioredis.from_url(settings.REDIS_URL, decode_responses=True)


class BulkService:
    """Bulk import, export, and AI generation operations."""

    def __init__(self, db: AsyncSession):
        self.db = db
        self.product_repo = ProductRepository(db)

    # ------------------------------------------------------------------
    # CSV / Excel Import
    # ------------------------------------------------------------------

    async def parse_csv(self, content: bytes) -> list[dict]:
        """Parse CSV content into a list of product dicts."""
        text = content.decode("utf-8-sig")  # Handle BOM
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        for i, row in enumerate(reader, start=2):
            cleaned = {}
            for key, value in row.items():
                key = key.strip().lower().replace(" ", "_")
                if key in TEMPLATE_COLUMNS and value and value.strip():
                    cleaned[key] = value.strip()
            if "sku" not in cleaned or "product_name" not in cleaned:
                raise ValidationError(
                    f"Row {i}: 'sku' and 'product_name' are required fields."
                )
            # Type conversions
            for num_field in ["price", "mrp", "cost_price", "weight"]:
                if num_field in cleaned:
                    try:
                        cleaned[num_field] = float(cleaned[num_field])
                    except ValueError:
                        cleaned[num_field] = None
            if "stock" in cleaned:
                try:
                    cleaned["stock"] = int(cleaned["stock"])
                except ValueError:
                    cleaned["stock"] = 0
            if "status" not in cleaned:
                cleaned["status"] = "draft"
            rows.append(cleaned)
        return rows

    async def parse_excel(self, content: bytes) -> list[dict]:
        """Parse Excel (xlsx) content into a list of product dicts."""
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValidationError(
                "Excel support requires 'openpyxl'. Install it with: pip install openpyxl"
            )

        wb = load_workbook(filename=io.BytesIO(content), read_only=True)
        ws = wb.active
        if not ws:
            raise ValidationError("Excel file has no active worksheet.")

        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValidationError("Excel file is empty.")

        headers = [
            str(h).strip().lower().replace(" ", "_") if h else f"col_{i}"
            for i, h in enumerate(header_row)
        ]

        rows = []
        for row_num, row in enumerate(rows_iter, start=2):
            cleaned = {}
            for header, value in zip(headers, row):
                if header in TEMPLATE_COLUMNS and value is not None and str(value).strip():
                    cleaned[header] = str(value).strip()
            if "sku" not in cleaned or "product_name" not in cleaned:
                continue  # Skip incomplete rows in Excel
            for num_field in ["price", "mrp", "cost_price", "weight"]:
                if num_field in cleaned:
                    try:
                        cleaned[num_field] = float(cleaned[num_field])
                    except ValueError:
                        cleaned[num_field] = None
            if "stock" in cleaned:
                try:
                    cleaned["stock"] = int(cleaned["stock"])
                except ValueError:
                    cleaned["stock"] = 0
            if "status" not in cleaned:
                cleaned["status"] = "draft"
            rows.append(cleaned)

        wb.close()
        return rows

    async def bulk_create_products(
        self, user_id: uuid.UUID, rows: list[dict], job_id: str
    ) -> dict:
        """Create products from parsed rows and track progress in Redis."""
        redis = _get_redis()
        total = len(rows)
        created = 0
        errors = []

        await redis.hset(f"bulk_job:{job_id}", mapping={
            "status": "processing",
            "total": str(total),
            "completed": "0",
            "errors": "0",
            "created": "0",
        })

        for i, row in enumerate(rows):
            try:
                row["user_id"] = user_id
                # Check SKU uniqueness
                if await self.product_repo.sku_exists(row["sku"]):
                    errors.append({"row": i + 2, "sku": row["sku"], "error": "SKU already exists"})
                else:
                    await self.product_repo.create(row)
                    created += 1
            except Exception as e:
                errors.append({"row": i + 2, "sku": row.get("sku", "?"), "error": str(e)})

            await redis.hset(f"bulk_job:{job_id}", mapping={
                "completed": str(i + 1),
                "created": str(created),
                "errors": str(len(errors)),
            })

        await redis.hset(f"bulk_job:{job_id}", mapping={
            "status": "completed",
            "error_details": json.dumps(errors[:50]),  # Cap error details
        })
        await redis.expire(f"bulk_job:{job_id}", 3600)  # TTL 1 hour
        await redis.aclose()

        return {
            "total": total,
            "created": created,
            "errors": len(errors),
            "error_details": errors[:20],
        }

    # ------------------------------------------------------------------
    # Job Status
    # ------------------------------------------------------------------

    async def get_job_status(self, job_id: str) -> dict:
        """Get the status of a bulk import/generate job."""
        redis = _get_redis()
        data = await redis.hgetall(f"bulk_job:{job_id}")
        await redis.aclose()

        if not data:
            return {"status": "not_found", "job_id": job_id}

        result = {"job_id": job_id}
        for key in ["status", "total", "completed", "created", "errors"]:
            if key in data:
                result[key] = int(data[key]) if key != "status" else data[key]
        if "error_details" in data:
            try:
                result["error_details"] = json.loads(data["error_details"])
            except json.JSONDecodeError:
                result["error_details"] = []
        return result

    # ------------------------------------------------------------------
    # Bulk AI Generate
    # ------------------------------------------------------------------

    async def bulk_generate_listings(
        self,
        user_id: uuid.UUID,
        product_ids: list[uuid.UUID],
        marketplaces: list[str],
        tone: str,
        provider: str | None,
        job_id: str,
    ) -> dict:
        """Generate AI listings for multiple products across marketplaces."""
        from app.services.listing_generator_service import ListingGeneratorService

        redis = _get_redis()
        total = len(product_ids) * len(marketplaces)
        completed = 0
        errors = []

        await redis.hset(f"bulk_job:{job_id}", mapping={
            "status": "processing",
            "total": str(total),
            "completed": "0",
            "created": "0",
            "errors": "0",
        })

        gen_service = ListingGeneratorService(self.db)
        created = 0

        for pid in product_ids:
            for mp in marketplaces:
                try:
                    await gen_service.generate_listing(
                        user_id=user_id,
                        product_id=pid,
                        marketplace=mp,
                        tone=tone,
                        provider_name=provider,
                    )
                    created += 1
                except Exception as e:
                    errors.append({"product_id": str(pid), "marketplace": mp, "error": str(e)})

                completed += 1
                await redis.hset(f"bulk_job:{job_id}", mapping={
                    "completed": str(completed),
                    "created": str(created),
                    "errors": str(len(errors)),
                })

        await redis.hset(f"bulk_job:{job_id}", mapping={
            "status": "completed",
            "error_details": json.dumps(errors[:50]),
        })
        await redis.expire(f"bulk_job:{job_id}", 3600)
        await redis.aclose()

        return {
            "total": total,
            "created": created,
            "errors": len(errors),
            "error_details": errors[:20],
        }

    # ------------------------------------------------------------------
    # Export
    # ------------------------------------------------------------------

    async def export_products_csv(self, user_id: uuid.UUID) -> str:
        """Export all user products as CSV string."""
        from sqlalchemy import select

        result = await self.db.execute(
            select(Product)
            .where(Product.user_id == user_id, Product.deleted_at.is_(None))
            .order_by(Product.created_at.desc())
        )
        products = result.scalars().all()

        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=TEMPLATE_COLUMNS)
        writer.writeheader()

        for p in products:
            row = {}
            for col in TEMPLATE_COLUMNS:
                val = getattr(p, col, None)
                row[col] = str(val) if val is not None else ""
            writer.writerow(row)

        return output.getvalue()

    def generate_template_csv(self) -> str:
        """Generate a blank CSV template with headers."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(TEMPLATE_COLUMNS)
        # Add an example row
        writer.writerow([
            "SKU-001", "Example Product Name", "BrandName", "Electronics",
            "A detailed product description", "Short desc", "999.00", "1299.00",
            "500.00", "100", "Black", "Medium", "Plastic", "0.5", "India",
            "Manufacturer Co.", "123 Street, City", "Packer Co.", "Importer Co.",
            "draft",
        ])
        return output.getvalue()
